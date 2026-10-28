"""Excel workbooks: listening history and library (favorite albums, playlists).

Used by the app (download buttons) and by the sync script, which keeps
up-to-date copies in the "Exports Excel" folder after each local sync.
"""

from __future__ import annotations

import io
import logging
import os
from pathlib import Path

import pandas as pd
from sqlalchemy.engine import Engine

from core import stats

log = logging.getLogger(__name__)

HISTORY_FILE = "Historique des écoutes.xlsx"
LIBRARY_FILE = "Bibliothèque - albums et playlists.xlsx"

LABELS = {
    "fr": {
        "sheet_plays": "Écoutes", "sheet_tracks": "Titres", "sheet_albums": "Albums",
        "sheet_artists": "Artistes", "sheet_days": "Par jour", "sheet_fav": "Albums favoris",
        "sheet_playlists": "Playlists", "sheet_pl_tracks": "Titres des playlists",
        "date": "Date (heure locale)", "date_utc": "Date (UTC)", "track": "Titre", "artist": "Artiste",
        "album": "Album", "genre": "Genre", "listened": "Écouté (s)", "estimated": "Durée estimée",
        "plays": "Écoutes", "minutes": "Minutes", "first": "Première écoute", "last": "Dernière écoute",
        "day": "Jour", "release": "Date de sortie", "year": "Année", "label": "Label", "added": "Ajouté le",
        "playlist": "Playlist", "creator": "Créée par", "own": "Ma playlist", "nb_tracks": "Titres",
        "duration": "Durée (min)", "artists": "Artistes", "played_tracks": "Titres écoutés",
        "played_share": "Part écoutée", "position": "Position", "track_id": "ID titre",
        "album_id": "ID album", "playlist_id": "ID playlist", "yes": "oui", "no": "non",
    },
    "en": {
        "sheet_plays": "Plays", "sheet_tracks": "Tracks", "sheet_albums": "Albums",
        "sheet_artists": "Artists", "sheet_days": "Per day", "sheet_fav": "Favorite albums",
        "sheet_playlists": "Playlists", "sheet_pl_tracks": "Playlist tracks",
        "date": "Date (local time)", "date_utc": "Date (UTC)", "track": "Track", "artist": "Artist",
        "album": "Album", "genre": "Genre", "listened": "Listened (s)", "estimated": "Estimated duration",
        "plays": "Plays", "minutes": "Minutes", "first": "First play", "last": "Last play",
        "day": "Day", "release": "Release date", "year": "Year", "label": "Label", "added": "Added on",
        "playlist": "Playlist", "creator": "Created by", "own": "My playlist", "nb_tracks": "Tracks",
        "duration": "Duration (min)", "artists": "Artists", "played_tracks": "Tracks played",
        "played_share": "Share played", "position": "Position", "track_id": "Track ID",
        "album_id": "Album ID", "playlist_id": "Playlist ID", "yes": "yes", "no": "no",
    },
}


def _naive(s: pd.Series) -> pd.Series:
    """Excel cannot store timezones: keep the local wall-clock time."""
    s = pd.to_datetime(s, errors="coerce")
    return s.dt.tz_localize(None) if getattr(s.dt, "tz", None) is not None else s


def history_sheets(plays: pd.DataFrame, lang: str = "fr") -> dict[str, pd.DataFrame]:
    L = LABELS[lang]
    yes_no = {True: L["yes"], False: L["no"]}
    p = plays.sort_values("played_at", ascending=False)
    sheets = {L["sheet_plays"]: pd.DataFrame({
        L["date"]: _naive(p["played_local"]), L["date_utc"]: _naive(p["played_at"]),
        L["track"]: p["track_title"], L["artist"]: p["artist_name"], L["album"]: p["album_title"],
        L["genre"]: p["genre"], L["listened"]: p["listened_seconds"].astype(int),
        L["estimated"]: p["duration_is_estimated"].map(yes_no), L["track_id"]: p["track_id"],
        L["album_id"]: p["album_id"],
    })}
    if p.empty:
        return sheets

    def agg(keys: list[str], cols: dict[str, str]) -> pd.DataFrame:
        g = p.groupby(keys, dropna=False).agg(plays=("play_id", "count"), minutes=("minutes", "sum"),
                                               first=("played_local", "min"), last=("played_local", "max"))
        g = g.reset_index().sort_values(["plays", "minutes"], ascending=False)
        g["minutes"] = g["minutes"].round(1)
        g["first"], g["last"] = _naive(g["first"]), _naive(g["last"])
        return g.rename(columns={**cols, "plays": L["plays"], "minutes": L["minutes"],
                                 "first": L["first"], "last": L["last"]})

    sheets[L["sheet_tracks"]] = agg(["track_title", "artist_name", "album_title"],
                                    {"track_title": L["track"], "artist_name": L["artist"], "album_title": L["album"]})
    sheets[L["sheet_albums"]] = agg(["album_title", "artist_name"], {"album_title": L["album"], "artist_name": L["artist"]})
    sheets[L["sheet_artists"]] = agg(["artist_name"], {"artist_name": L["artist"]})
    days = p.groupby(p["played_local"].dt.date).agg(plays=("play_id", "count"), minutes=("minutes", "sum"))
    days = days.reset_index().rename(columns={"played_local": L["day"], "plays": L["plays"], "minutes": L["minutes"]})
    days[L["minutes"]] = days[L["minutes"]].round(1)
    sheets[L["sheet_days"]] = days.sort_values(L["day"], ascending=False)
    return sheets


def library_sheets(favorites: pd.DataFrame, playlists: pd.DataFrame, playlist_tracks: pd.DataFrame,
                   lang: str = "fr") -> dict[str, pd.DataFrame]:
    L = LABELS[lang]
    yes_no = {True: L["yes"], False: L["no"]}
    fav = favorites
    sheets = {L["sheet_fav"]: pd.DataFrame({
        L["album"]: fav.get("title"), L["artist"]: fav.get("artist_name"),
        L["release"]: pd.to_datetime(fav.get("release_date"), errors="coerce").dt.date
        if not fav.empty else fav.get("release_date"),
        L["genre"]: fav.get("genre"), L["label"]: fav.get("label"), L["added"]: _naive(fav.get("added_at"))
        if not fav.empty else fav.get("added_at"),
        L["plays"]: fav.get("plays"), L["minutes"]: fav.get("minutes", pd.Series(dtype=float)).round(1),
        L["album_id"]: fav.get("album_id"),
    })}
    pl = playlists
    sheets[L["sheet_playlists"]] = pd.DataFrame({
        L["playlist"]: pl.get("title"), L["creator"]: pl.get("creator_name"),
        L["own"]: pl.get("is_own", pd.Series(dtype=bool)).map(yes_no), L["nb_tracks"]: pl.get("tracks"),
        L["duration"]: (pl.get("duration", pd.Series(dtype=float)) / 60).round(0),
        L["artists"]: pl.get("artists"), L["played_tracks"]: pl.get("played_tracks"),
        L["played_share"]: pl.get("played_share", pd.Series(dtype=float)).round(3), L["plays"]: pl.get("plays"),
        L["added"]: _naive(pl.get("added_at")) if not pl.empty else pl.get("added_at"),
        L["playlist_id"]: pl.get("id"),
    })
    titles = dict(zip(pl.get("id", []), pl.get("title", [])))
    pt = playlist_tracks
    sheets[L["sheet_pl_tracks"]] = pd.DataFrame({
        L["playlist"]: pt.get("playlist_id", pd.Series(dtype=int)).map(titles), L["position"]: pt.get("position", pd.Series(dtype=int)) + 1,
        L["track"]: pt.get("track_title"), L["artist"]: pt.get("artist_name"), L["album"]: pt.get("album_title"),
        L["plays"]: pt.get("plays"), L["added"]: _naive(pt.get("added_at")) if not pt.empty else pt.get("added_at"),
        L["track_id"]: pt.get("track_id"),
    })
    return sheets


def to_xlsx(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Workbook with a frozen, filterable header and readable column widths."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2A78D6")
            for i, col in enumerate(df.columns, start=1):
                sample = df[col].head(500).astype(str).map(len)
                width = max([len(str(col))] + ([int(sample.max())] if len(sample) else []))
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
                if pd.api.types.is_datetime64_any_dtype(df[col]):
                    for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                        for c in cell:
                            c.number_format = "dd/mm/yyyy hh:mm"
    return buf.getvalue()


def _atomic_write(path: Path, data: bytes) -> None:
    tmp = path.with_suffix(".tmp.xlsx")
    tmp.write_bytes(data)
    os.replace(tmp, path)


def write_exports(engine: Engine, out_dir: Path, tz: str, lang: str = "fr") -> list[Path]:
    """Rewrite the Excel copies of the whole history and library."""
    out_dir.mkdir(parents=True, exist_ok=True)
    raw = stats.load_raw(engine)
    plays = stats.assemble(raw, tz)
    paths = [out_dir / HISTORY_FILE, out_dir / LIBRARY_FILE]
    _atomic_write(paths[0], to_xlsx(history_sheets(plays, lang)))
    lib = stats.load_library(engine)
    genres = stats.album_genres(raw.albums, raw.tags, raw.aliases)
    _atomic_write(paths[1], to_xlsx(library_sheets(
        stats.favorite_albums_view(lib, plays, genres), stats.playlists_view(lib, plays),
        stats.playlist_tracks_view(lib, plays), lang)))
    log.info("Excel exports written to %s", out_dir)
    return paths
