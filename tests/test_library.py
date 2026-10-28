import io
from datetime import datetime

import openpyxl
import pandas as pd
import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from core import stats
from core.deezer_export import import_rows, match_columns, parse_rows
from core.excel import history_sheets, library_sheets, to_xlsx
from core.library import store_favorite_albums, store_playlist_list, store_playlist_tracks
from core.models import Album, FavoriteAlbum, Play, Playlist, PlaylistTrack
from core.poster import PAGE_H, PAGE_W, album_poster_pdf, best_layout
from core.sync import run_sync
from core.deezer_gw import parse_history

FAV = [
    {"ALB_ID": "10", "ALB_TITLE": "Album A", "ART_ID": "1", "ART_NAME": "Alpha", "ALB_PICTURE": "abc",
     "DATE_FAVORITE": "2026-10-08 05:33:30", "PHYSICAL_RELEASE_DATE": "2008-06-17"},
    {"ALB_ID": "20", "ALB_TITLE": "Album B", "ART_ID": "2", "ART_NAME": "Beta", "ALB_PICTURE": "",
     "DATE_FAVORITE": "2025-01-01 10:00:00", "PHYSICAL_RELEASE_DATE": "0000-00-00"},
]
PLAYLISTS = [
    {"PLAYLIST_ID": "500", "TITLE": "Mine", "PARENT_USER_ID": "42", "PARENT_USERNAME": "Me", "NB_SONG": 2,
     "DATE_ADD": "2026-01-01 00:00:00", "PLAYLIST_PICTURE": "p1", "PICTURE_TYPE": "playlist"},
    {"PLAYLIST_ID": "600", "TITLE": "Theirs", "PARENT_USER_ID": "7", "PARENT_USERNAME": "Other", "NB_SONG": 1,
     "DATE_FAVORITE": "2026-02-01 00:00:00"},
]
SONG = lambda i: {"SNG_ID": str(i), "SNG_TITLE": f"Song {i}", "ART_ID": "1", "ART_NAME": "Alpha", "ALB_ID": "10",  # noqa: E731
                  "ALB_TITLE": "Album A", "DURATION": "200", "DATE_ADD": 1700000000}


def test_library_store_and_views(engine, history_response):
    run_sync(engine, type("S", (), {"fetch_recent": lambda self: parse_history(history_response)})(), enrich=False)
    with Session(engine) as s:
        assert store_favorite_albums(s, FAV, 42) == 2
        to_fetch = store_playlist_list(s, PLAYLISTS, 42)
        assert sorted(to_fetch) == [500, 600]
        store_playlist_tracks(s, 500, {"CHECKSUM": "x", "DURATION": "400"}, [SONG(67238732), SONG(99)])
        store_playlist_tracks(s, 600, {}, [SONG(99)])
        s.commit()
        assert s.get(Album, 10).release_date == "2008-06-17"
        assert s.get(Album, 20).release_date is None
        assert s.get(Playlist, 500).is_own and not s.get(Playlist, 600).is_own
        # Refresh: unchanged playlists are not reloaded, removed ones disappear.
        assert store_playlist_list(s, PLAYLISTS[:1], 42) == []
        s.commit()
        assert s.scalar(select(func.count()).select_from(Playlist)) == 1
        assert s.scalar(select(func.count()).select_from(PlaylistTrack)) == 2
        store_favorite_albums(s, FAV[:1], 42)
        s.commit()
        assert s.scalar(select(func.count()).select_from(FavoriteAlbum)) == 1

    plays = stats.assemble(stats.load_raw(engine))
    lib = stats.load_library(engine)
    fav = stats.favorite_albums_view(lib, plays)
    assert fav.iloc[0]["release_year"] == 2008
    pl = stats.playlists_view(lib, plays)
    row = pl.iloc[0]
    assert row["tracks"] == 2 and row["played_tracks"] == 1 and row["played_share"] == 0.5

    book = openpyxl.load_workbook(io.BytesIO(to_xlsx(library_sheets(
        fav, pl, stats.playlist_tracks_view(lib, plays)))))
    assert book.sheetnames == ["Albums favoris", "Playlists", "Titres des playlists"]
    assert book["Titres des playlists"].max_row == 3


def test_history_workbook(engine, history_response):
    run_sync(engine, type("S", (), {"fetch_recent": lambda self: parse_history(history_response)})(), enrich=False)
    plays = stats.assemble(stats.load_raw(engine))
    for lang, first in (("fr", "Écoutes"), ("en", "Plays")):
        book = openpyxl.load_workbook(io.BytesIO(to_xlsx(history_sheets(plays, lang))))
        assert book.sheetnames[0] == first
        assert book[first].max_row == 7  # header + 6 plays
    assert to_xlsx(history_sheets(plays.iloc[0:0]))  # empty history still gives a workbook


@pytest.mark.parametrize("n", [1, 7, 50, 336, 1000])
def test_poster_layout_fits_page(n):
    lay = best_layout(n)
    assert lay.cols * lay.rows >= n
    assert lay.cols * lay.cell_w <= PAGE_W and lay.rows * lay.cell_h <= PAGE_H
    assert lay.cover <= lay.cell_w and lay.cover * 1.3 <= lay.cell_h


def test_poster_pdf_without_network():
    pdf = album_poster_pdf([{"title": "Été 85 ✨", "artist": "Ø", "release": "1985", "cover_url": None}] * 5, "Test")
    assert pdf.startswith(b"%PDF") and len(pdf) > 1000


def test_export_columns_and_rows():
    df = pd.DataFrame({
        "Song Title": ["Instant Crush", "Within", None], "Artist": ["Daft Punk"] * 3,
        "ISRC": ["USQX91300105", "", "X"], "Album Title": ["RAM"] * 3,
        "Listening Time": [120, "45", 3], "Date": ["2025-03-01 10:00:00", "2025-03-01 10:05:00", "bad"],
        "IP Address": ["1.2.3.4"] * 3,
    })
    mapping = match_columns(df.columns)
    assert mapping == {"title": "Song Title", "artist": "Artist", "isrc": "ISRC", "album": "Album Title",
                       "listened": "Listening Time", "date": "Date"}
    rows = parse_rows(df, mapping, tz="Europe/Zurich")
    assert len(rows) == 2
    assert rows[0].played_at == datetime(2025, 3, 1, 9, 0) and rows[0].listened == 120
    assert rows[1].isrc is None


class FakePublic:
    def get(self, path, retries=4):
        if path == "track/isrc:USQX91300105":
            return {"id": 67238732, "title": "Instant Crush", "isrc": "USQX91300105", "duration": 337,
                    "artist": {"id": 27, "name": "Daft Punk"}, "album": {"id": 6575789, "title": "RAM"}}
        return None

    def track(self, tid):
        return None


def test_import_skips_duplicates(engine, history_response):
    from core.deezer_export import ExportRow

    run_sync(engine, type("S", (), {"fetch_recent": lambda self: parse_history(history_response)})(), enrich=False)
    rows = [
        ExportRow(datetime(2026, 10, 28, 12, 3, 0), "Instant Crush", "Daft Punk", None, "USQX91300105", 300),  # dup
        ExportRow(datetime(2025, 1, 1, 8, 0, 0), "Instant Crush", "Daft Punk", None, "USQX91300105", 30),
        ExportRow(datetime(2025, 1, 1, 9, 0, 0), "Unknown", "Nobody", None, None, 10),
    ]
    with Session(engine) as s:
        report = import_rows(s, rows, client=FakePublic())
        s.commit()
        assert (report.inserted, report.duplicates, report.unresolved) == (1, 1, 1)
        imported = s.scalars(select(Play).where(Play.source == "deezer_export")).one()
        assert imported.listened_seconds == 30 and not imported.duration_is_estimated
